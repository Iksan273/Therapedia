# -*- coding: utf-8 -*-
"""
Generator dokumen Final Requirement (SRS) Therapedia -> .docx
Jalankan:  python docs/generate_requirement.py
Revisi konten cukup di bagian DATA di bawah, lalu jalankan ulang.
"""
import os
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------------------------------------------------------------------------
# META
# ---------------------------------------------------------------------------
VERSION = "1.0"
DOC_DATE = date.today().strftime("%d/%m/%Y")
PROJECT = "One Gate Integrated Clinic System"
CLIENT = "Therapedia Center (Pediatric Developmental Center)"
VENDOR = "Jupiter Cosmic (IT Solutions)"
OUT_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        f"Therapedia_Final_Requirement_v{VERSION}.docx")

BRAND = "007AFF"
NAVY = RGBColor(0x04, 0x0E, 0x1E)
BLUE = RGBColor(0x00, 0x7A, 0xFF)

# ---------------------------------------------------------------------------
# HELPERS
# ---------------------------------------------------------------------------
doc = Document()


def set_cell_bg(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def add_field(paragraph, instr):
    """Sisipkan field Word (PAGE, TOC, dll)."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Klik kanan > Update Field untuk menampilkan daftar isi." if "TOC" in instr else "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, text, sep, placeholder, end):
        run._r.append(el)


def h(text, level=1):
    return doc.add_heading(text, level=level)


def p(text="", bold=False, italic=False, size=None, align=None, color=None, space_after=None):
    para = doc.add_paragraph()
    if text:
        run = para.add_run(text)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
    if align is not None:
        para.alignment = align
    if space_after is not None:
        para.paragraph_format.space_after = Pt(space_after)
    return para


def rich(parts, style=None):
    """parts: list of (text, bold)."""
    para = doc.add_paragraph(style=style)
    for text, bold in parts:
        r = para.add_run(text)
        r.bold = bold
    return para


def bullets(items, style="List Bullet"):
    for it in items:
        if isinstance(it, tuple):
            rich([(it[0], True), (it[1], False)], style=style)
        else:
            doc.add_paragraph(it, style=style)


def table(headers, rows, widths=None, font_size=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0]
    repeat_header(hdr)
    for i, text in enumerate(headers):
        c = hdr.cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(text)
        r.bold = True
        r.font.size = Pt(font_size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_bg(c, BRAND)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(text))
            r.font.size = Pt(font_size)
            if i == 0:
                r.bold = True
            if ri % 2 == 1:
                set_cell_bg(cells[i], "EEF5FF")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def fr_table(code_prefix, items):
    """items: list of (judul, deskripsi, prioritas)."""
    rows = [(f"{code_prefix}-{i:02d}", j, d, pr) for i, (j, d, pr) in enumerate(items, 1)]
    table(["ID", "Kebutuhan", "Deskripsi / Kriteria Penerimaan", "Prioritas"],
          rows, widths=[2.4, 3.6, 9.0, 1.8])


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
# STYLES & LAYOUT
# ---------------------------------------------------------------------------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
normal.paragraph_format.space_after = Pt(6)

for lvl, size in ((1, 16), (2, 13), (3, 11.5)):
    st = doc.styles[f"Heading {lvl}"]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = NAVY if lvl == 1 else BLUE
    rfonts = st.element.rPr.rFonts
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)  # pakai font eksplisit, bukan font tema

sec = doc.sections[0]
sec.orientation = WD_ORIENT.PORTRAIT
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.2)
sec.different_first_page_header_footer = True

hp = sec.header.paragraphs[0]
hp.text = f"Therapedia – Final Requirement v{VERSION} – Rahasia (Confidential)"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.runs[0].font.size = Pt(8.5)
hp.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run("Halaman ").font.size = Pt(9)
add_field(fp, "PAGE")
fp.add_run(f"   |   {VENDOR}").font.size = Pt(9)

# ---------------------------------------------------------------------------
# COVER
# ---------------------------------------------------------------------------
for _ in range(5):
    p()
p("DOKUMEN SPESIFIKASI KEBUTUHAN", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, color=BLUE)
p("FINAL REQUIREMENT", bold=True, size=28, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY, space_after=0)
p(PROJECT, bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY)
p("Sistem Informasi Terpadu Satu Pintu Operasional Klinik & Rekam Klinis Digital",
  italic=True, size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(4):
    p()
cover = doc.add_table(rows=0, cols=2)
cover.alignment = WD_TABLE_ALIGNMENT.CENTER
for k, v in (("Disiapkan untuk", CLIENT), ("Disiapkan oleh", VENDOR),
             ("Versi Dokumen", VERSION), ("Tanggal", DOC_DATE),
             ("Status", "Menunggu Persetujuan Klien")):
    row = cover.add_row().cells
    row[0].text, row[1].text = k, f": {v}"
    row[0].paragraphs[0].runs[0].bold = True
    row[0].width, row[1].width = Cm(4.5), Cm(9)
page_break()

# ---------------------------------------------------------------------------
# KONTROL DOKUMEN + DAFTAR ISI
# ---------------------------------------------------------------------------
h("Kontrol Dokumen", 1)
table(["Versi", "Tanggal", "Penyusun", "Keterangan"],
      [(VERSION, DOC_DATE, VENDOR, "Final requirement untuk review & persetujuan klien")],
      widths=[2, 3, 5, 7])
p("Daftar Isi", bold=True, size=14, color=NAVY)
add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
page_break()

# ---------------------------------------------------------------------------
# 1. PENDAHULUAN
# ---------------------------------------------------------------------------
h("1. Pendahuluan", 1)
h("1.1 Tujuan Dokumen", 2)
p("Dokumen ini merangkum seluruh kebutuhan fungsional dan non-fungsional final untuk "
  f"pengembangan {PROJECT} milik {CLIENT}. Dokumen ini menjadi acuan bersama antara "
  "Therapedia dan tim pengembang dalam proses pengembangan, pengujian (UAT), serta serah terima sistem. "
  "Setelah disetujui, setiap perubahan di luar dokumen ini diproses melalui mekanisme Change Request (CR).")

h("1.2 Ruang Lingkup Produk", 2)
p("Sistem merupakan aplikasi web terpadu (satu pintu) yang mencakup:")
bullets([
    ("Website publik (Landing Page) ", "– profil klinik, program, tim terapis, cabang, artikel, dan formulir kontak."),
    ("Alur Inquiry & Asesmen ", "– penerimaan calon klien, pemilihan layanan, kuesioner orang tua, jadwal asesmen, hingga keputusan admisi."),
    ("Alur Penjadwalan & Sesi Terapi ", "– kalender mingguan multi-terapis, sesi berulang, pembatalan, reschedule, dan laporan sesi."),
    ("Keuangan & Kredit Sesi ", "– paket terapi, invoice, verifikasi bukti transfer, kredit sesi, dan buku besar kredit."),
    ("Portal Terapis ", "– jadwal pribadi, laporan sesi, dan berkas klien."),
    ("Portal Orang Tua ", "– status tagihan, unggah bukti bayar, saldo kredit, dan riwayat laporan sesi anak."),
    ("Dashboard Eksekutif & Administrasi ", "– analitik pendapatan, performa cabang, manajemen pengguna, dan hak akses (RBAC)."),
])

h("1.3 Definisi & Istilah", 2)
table(["Istilah", "Definisi"], [
    ("Inquiry", "Calon klien yang baru menghubungi klinik / didaftarkan admin sebelum resmi menjadi klien aktif."),
    ("Intake", "Proses pendataan awal calon klien (data anak & orang tua, cabang, layanan)."),
    ("Pipeline", "Tahapan status calon klien dari Inquiry hingga keputusan akhir (Admitted / Done / Discontinued)."),
    ("Kode ASM", "Kode akses unik (format ASM-XXXX) untuk orang tua/guru mengisi kuesioner asesmen tanpa login."),
    ("Kode TDC", "Kode akses unik klien (format TDC-XXXX) untuk login orang tua ke Portal Orang Tua."),
    ("Paket / Kredit", "Paket terapi berisi sejumlah sesi; 1 kredit = 1 sesi terapi."),
    ("Frozen", "Status klien yang saldo kreditnya habis (0); sesi tetap terlihat namun ditandai belum terbayar."),
    ("Discharge", "Pengakhiran status klien aktif dengan alasan tertentu."),
    ("BOT-A / FOT-A", "Jenis layanan asesmen terapi okupasi yang dapat dipilih pada tahap intake."),
    ("SI / NDT", "Sensory Integration / Neurodevelopmental Treatment – pendekatan terapi klinik."),
    ("Kuadran Sensori", "Pengelompokan skor Winnie Dunn Sensory Profile: AV (Avoiding), SN (Sensitivity), RG (Registration), SK (Seeking)."),
    ("RBAC", "Role-Based Access Control – pengaturan hak akses modul berdasarkan peran pengguna."),
], widths=[3.5, 13.3])

# ---------------------------------------------------------------------------
# 2. GAMBARAN UMUM
# ---------------------------------------------------------------------------
h("2. Gambaran Umum Sistem", 1)
h("2.1 Cabang", 2)
table(["Kode", "Cabang", "Lokasi"], [
    ("EAST", "East Branch", "Ruko Rungkut Megah Raya, Surabaya"),
    ("WEST", "West Branch", "Lagoon Avenue Sungkono Mall, Surabaya"),
    ("CTL", "Citraland Branch", "Royal Park, Citraland, Lakarsantri, Surabaya"),
], widths=[2.2, 4, 10.6])
p("Seluruh data operasional (klien, jadwal, terapis, invoice) terikat pada satu cabang. "
  "Master Director dapat melihat data gabungan seluruh cabang maupun per cabang; peran lain dibatasi pada cabangnya.")

h("2.2 Alur Bisnis Utama", 2)
p("A. Alur Inquiry → Admisi", bold=True)
bullets([
    "Calon klien didaftarkan (New Intake) → status Inquiry.",
    "Admin memilih layanan → status Service Selected.",
    "Admin membuat kode kuesioner (ASM) & menjadwalkan sesi asesmen → status Assessment Scheduled.",
    "Orang tua mengisi kuesioner / sesi asesmen selesai → status Assessment Done.",
    "Invoice diterbitkan, orang tua mengunggah bukti transfer, Finance memverifikasi.",
    "Keputusan akhir: Admitted (menjadi klien aktif), Done Consult, Done Assessment, atau Discontinued (dengan alasan).",
], style="List Number")
p("B. Alur Penjadwalan → Sesi → Kredit", bold=True)
bullets([
    "Admin Schedule membuat jadwal tunggal atau berulang (mingguan) untuk klien aktif.",
    "Sesi dilaksanakan; Admin menandai Selesai → 1 kredit terpotong.",
    "Terapis mengisi laporan sesi (Activity, Note, Homework) yang dapat dilihat orang tua.",
    "Bila kredit habis, klien berstatus Frozen hingga paket diperpanjang (renewal) melalui Finance.",
], style="List Number")

h("2.3 Daftar Modul", 2)
table(["Kode", "Modul", "Pengguna Utama"], [
    ("AUTH", "Autentikasi & Akses", "Semua pengguna"),
    ("LND", "Landing Page (Website Publik)", "Publik / calon orang tua"),
    ("INQ", "Inquiry & Intake", "Admin Inquiry, Branch Manager, Master"),
    ("ASM", "Asesmen & Kuesioner", "Admin Inquiry, Orang Tua/Guru, Terapis"),
    ("SCH", "Penjadwalan & Sesi", "Admin Schedule, Branch Manager, Master"),
    ("CLT", "Klien Aktif", "Admin Schedule, Branch Manager, Finance, Master"),
    ("FIN", "Keuangan, Paket & Kredit", "Finance, Master"),
    ("THR", "Portal Terapis", "Terapis"),
    ("PRT", "Portal Orang Tua", "Orang Tua"),
    ("EXE", "Dashboard Eksekutif", "Master, Branch Manager, Finance"),
    ("ADM", "Administrasi Pengguna & RBAC", "Master"),
], widths=[2, 6.5, 8.3])

# ---------------------------------------------------------------------------
# 3. PERAN
# ---------------------------------------------------------------------------
h("3. Peran Pengguna & Hak Akses", 1)
h("3.1 Deskripsi Peran", 2)
table(["Peran", "Cakupan Data", "Tanggung Jawab Utama"], [
    ("Master Director (HQ)", "Seluruh cabang (dapat memilih All / per cabang)",
     "Akses penuh seluruh modul, analitik pendapatan & performa cabang, kelola pengguna dan RBAC."),
    ("Branch Manager", "Cabang sendiri",
     "Memantau pendapatan cabang, inquiry, jadwal, dan klien aktif cabang."),
    ("Admin Inquiry", "Cabang sendiri",
     "Mengelola intake, pipeline, kuesioner asesmen, dan keputusan admisi."),
    ("Admin Schedule", "Cabang sendiri",
     "Mengelola kalender, jadwal sesi, penyelesaian/pembatalan/reschedule sesi, dan klien aktif."),
    ("Finance", "Sesuai penugasan",
     "Verifikasi bukti transfer, penerbitan invoice, renewal kredit, master paket, audit kredit."),
    ("Terapis", "Klien yang ditangani",
     "Melihat jadwal pribadi, mengisi laporan sesi, melihat berkas klien & hasil kuesioner (read-only)."),
    ("Orang Tua (Client)", "Data anak sendiri",
     "Melihat tagihan, mengunggah bukti bayar, melihat saldo kredit & laporan sesi anak."),
], widths=[3.6, 4.2, 9])

h("3.2 Matriks Hak Akses Default", 2)
p("Matriks berikut adalah konfigurasi awal. Master Director dapat mengubahnya melalui modul RBAC, "
  "termasuk membuat peran baru (custom role).")
Y, N = "✓", "–"
table(["Modul", "Master", "Manager", "Adm. Inquiry", "Adm. Schedule", "Finance", "Terapis"], [
    ("Dashboard Revenue", Y, Y, N, N, Y, N),
    ("Inquiry Pipeline", Y, Y, Y, N, N, N),
    ("Inquiry Dashboard", Y, Y, Y, N, N, N),
    ("Weekly Calendar", Y, Y, N, Y, N, N),
    ("Schedule Dashboard", Y, Y, N, Y, N, N),
    ("Active Client", Y, Y, N, Y, Y, N),
    ("Finance & Invoices", Y, N, N, N, Y, N),
    ("User Management", Y, N, N, N, N, N),
    ("RBAC", Y, N, N, N, N, N),
    ("Therapist Module", Y, N, N, N, N, Y),
], widths=[4.4, 2, 2, 2.2, 2.4, 2, 1.8])
p("Portal Orang Tua bersifat terpisah dan hanya dapat diakses dengan Kode TDC milik klien.", italic=True)

# ---------------------------------------------------------------------------
# 4. KEBUTUHAN FUNGSIONAL
# ---------------------------------------------------------------------------
h("4. Kebutuhan Fungsional", 1)
p("Prioritas: Must = wajib pada rilis pertama; Should = sangat dianjurkan; Could = opsional bila waktu memungkinkan.",
  italic=True)

h("4.1 Autentikasi & Akses (AUTH)", 2)
fr_table("FR-AUTH", [
    ("Login staf", "Staf login menggunakan email dan kata sandi. Kata sandi divalidasi di server dan disimpan dalam bentuk hash. Login gagal menampilkan pesan error tanpa membocorkan detail akun.", "Must"),
    ("Login orang tua", "Orang tua login ke Portal Orang Tua menggunakan Kode TDC milik anak.", "Must"),
    ("Akses kuesioner", "Orang tua/guru membuka formulir kuesioner menggunakan Kode ASM tanpa perlu akun.", "Must"),
    ("Pengalihan sesuai peran", "Setelah login, pengguna diarahkan ke halaman utama sesuai perannya; halaman di luar haknya tidak dapat diakses (UI dan API).", "Must"),
    ("Pembatasan cabang", "Pengguna non-Master hanya melihat data cabangnya. Master dapat memilih All Branches atau cabang tertentu dari top bar.", "Must"),
    ("Logout & sesi", "Pengguna dapat logout; sesi berakhir otomatis setelah periode tidak aktif tertentu.", "Must"),
    ("Lupa kata sandi", "Staf dapat meminta reset kata sandi (melalui Master/admin atau tautan reset).", "Should"),
    ("Remember me", "Opsi tetap masuk pada perangkat tepercaya.", "Could"),
])

h("4.2 Landing Page / Website Publik (LND)", 2)
fr_table("FR-LND", [
    ("Halaman profil", "Menampilkan Hero, Why Us (4 pilar, tab SI vs NDT), Program (Regular OT, Sensory Spark, EIBI) beserta detail, Tim terapis (filter & profil), Cabang (foto, jam operasional, tautan Google Maps), Artikel/Knowledge, Kontak, dan Footer.", "Must"),
    ("Navigasi", "Header dengan navigasi antar-section (scroll-spy) dan tampilan responsif di perangkat mobile.", "Must"),
    ("Formulir kontak", "Field: Nama Orang Tua*, No. WhatsApp*, Email, Nama Anak, Usia Anak, Cabang*, Keluhan/Kebutuhan*. Data tersimpan di sistem sebagai lead dan dapat ditindaklanjuti Admin Inquiry.", "Must"),
    ("Chat WhatsApp", "Tombol membuka WhatsApp resmi klinik (wa.me) dengan pesan terformat.", "Must"),
    ("Pengelolaan konten", "Konten program, tim, cabang, dan artikel dapat diperbarui tanpa perubahan kode (CMS sederhana).", "Could"),
])

h("4.3 Inquiry & Intake (INQ)", 2)
fr_table("FR-INQ", [
    ("Pipeline Kanban", "Papan Kanban 8 kolom: Inquiry, Service Selected, Assessment Scheduled, Assessment Done, Admitted, Done Consult, Done Assessment, Discontinued. Mendukung pencarian (nama anak, orang tua, email, kode) dan filter cabang (Master).", "Must"),
    ("New Intake", "Formulir: Nama Anak*, Jenis Kelamin, Tanggal Lahir*, Nama Orang Tua*, Kontak Orang Tua*, Email Orang Tua*, Cabang*. Sistem otomatis membuat Kode TDC unik.", "Must"),
    ("Detail intake", "Halaman detail calon klien dengan langkah-langkah yang dapat diisi tidak berurutan: (1) data intake (dapat diedit), (2) pilihan layanan, (3) kode kuesioner, (4) jadwal asesmen, (5) hasil kuesioner, (6) link Google Drive, (7) invoice & bukti transfer, (8) keputusan akhir.", "Must"),
    ("Pemilihan layanan", "Dapat memilih lebih dari satu layanan: BOT-A, FOT-A, Konsultasi tanpa laporan, Konsultasi dengan laporan tertulis; plus opsi School Companion Profile (hanya untuk BOT-A/FOT-A). Status berubah menjadi Service Selected.", "Must"),
    ("Kode kuesioner", "Admin memilih kategori kuesioner dan membuat Kode ASM unik yang dapat disalin dan dikirim via WhatsApp. Status menjadi Assessment Scheduled.", "Must"),
    ("Jadwal asesmen", "Admin menjadwalkan sesi asesmen langsung dari halaman detail (menggunakan formulir jadwal pada modul SCH).", "Must"),
    ("Link Google Drive", "Admin menyimpan tautan folder Google Drive klien.", "Must"),
    ("Keputusan akhir", "Admit (menetapkan tanggal bergabung & menjadikan klien aktif), Done Consult, Done Assessment, atau Discontinue (alasan wajib diisi).", "Must"),
    ("Dashboard inquiry", "Filter cabang, periode (7 hari, bulan ini, bulan lalu, kuartal, custom), layanan, dan status (termasuk 'Inquiry Baru'). Grafik funnel pipeline + % konversi, distribusi layanan, tren intake 6 bulan.", "Must"),
    ("Daftar tindak lanjut", "Tab: kuesioner belum diisi (dengan tombol pengingat WhatsApp), daftar intake terfilter, dan log Discontinued.", "Must"),
    ("Riwayat pipeline", "Setiap perpindahan status tercatat (status lama/baru, waktu, pengguna) dan dapat dilihat pada detail klien.", "Should"),
])

h("4.4 Asesmen & Kuesioner (ASM)", 2)
fr_table("FR-ASM", [
    ("Master kategori", "Admin membuat, mengubah, dan menghapus kategori kuesioner (mis. Child Sensory Profile 2, School Companion, Fine Motor/VMI, Speech/Language, ADL).", "Must"),
    ("Domain & pertanyaan", "Setiap kategori terdiri dari domain (judul & teks pengantar) dan pertanyaan dengan 6 tipe: skala 0–5 (Winnie Dunn), rentang, pilihan ganda, checkbox multi, teks bebas, ya/tidak. Pertanyaan dapat dipetakan ke kuadran sensori (AV/SN/RG/SK) dan diduplikasi.", "Must"),
    ("Generate kode", "Kode ASM dibuat untuk klien & kategori tertentu, dari halaman intake maupun master data.", "Must"),
    ("Pengisian kuesioner", "Orang tua/guru memasukkan Kode ASM → mengisi formulir dengan indikator progres (%), filter item belum terjawab, dan tombol lompat ke item belum terjawab. Semua item wajib dijawab sebelum dikirim.", "Must"),
    ("Penyimpanan jawaban", "Jawaban tersimpan ke data klien; status pipeline otomatis menjadi Assessment Done. Halaman selesai menampilkan tautan konfirmasi WhatsApp.", "Must"),
    ("Tabel hasil", "Tampilan hasil dengan mode klinis / matriks / inquiry, filter domain, skor mentah per domain, total skor kuadran, dan usia anak terperinci (tahun, bulan, hari). Item yang tidak terjawab ditampilkan kosong (tidak diberi nilai default).", "Must"),
    ("Cetak hasil", "Tabel hasil dapat dicetak / disimpan sebagai PDF.", "Must"),
    ("Kode sekali pakai", "Kode ASM memiliki status (aktif/terpakai) dan opsional masa berlaku.", "Should"),
])

h("4.5 Penjadwalan & Sesi (SCH)", 2)
fr_table("FR-SCH", [
    ("Kalender mingguan", "Grid Senin–Sabtu pukul 08.00–17.00 dan tampilan agenda harian; navigasi minggu sebelumnya/hari ini/berikutnya; filter terapis; opsi tampilkan klien discharged; legenda status.", "Must"),
    ("Tambah jadwal tunggal", "Field: Klien, Layanan, Terapis, Paket kredit (atau 0 kredit/Frozen), Tanggal, Jam mulai & selesai, Catatan.", "Must"),
    ("Jadwal berulang", "Mode mingguan multi-hari: pilih hari, jam/terapis/layanan/paket per hari, berulang 4/8/12/16/24 minggu.", "Must"),
    ("Deteksi konflik", "Sistem memperingatkan bila jadwal di luar jam kerja terapis atau bentrok dengan jadwal lain (double booking).", "Must"),
    ("Penanda Frozen", "Sesi milik klien dengan saldo kredit 0 ditandai Frozen pada kalender.", "Must"),
    ("Detail sesi", "Melihat detail sesi dan laporan (Activity/Note/Homework).", "Must"),
    ("Selesaikan sesi", "Menandai sesi selesai: sesi terapi memotong 1 kredit; sesi asesmen mengubah status klien menjadi Assessment Done.", "Must"),
    ("Batalkan sesi", "Alasan: Sakit, Izin keluarga, Bentrok sekolah, Tanpa kabar, Lainnya + catatan. Potongan kredit mengikuti aturan BR-03.", "Must"),
    ("Reschedule", "Mengubah tanggal, jam, dan terapis. Reschedule ditolak bila terjadi konflik jadwal.", "Must"),
    ("Aksi massal (bulk)", "Memilih beberapa sesi lalu: Selesaikan massal, Reschedule massal (geser hari atau tanggal pasti, opsional ganti terapis), Batalkan massal (alasan + catatan) dengan aturan kredit yang sama seperti aksi tunggal.", "Must"),
    ("Dashboard jadwal", "Kartu statistik (klien aktif, terjadwal, selesai, dibatalkan); filter klien, terapis, cabang, status, rentang tanggal; daftar sesi; grafik jenis layanan & alasan pembatalan.", "Must"),
    ("Birthday radar", "Daftar klien berulang tahun bulan berjalan dengan tombol ucapan via WhatsApp.", "Should"),
])

h("4.6 Klien Aktif (CLT)", 2)
fr_table("FR-CLT", [
    ("Roster klien", "Tabel berhalaman berisi profil, orang tua, cabang, paket kredit, dan riwayat pembatalan; filter status kredit.", "Must"),
    ("Birthday hub", "Pemilih bulan dan tombol ucapan WhatsApp.", "Should"),
    ("Analitik klien", "Grafik kehadiran, beban kasus terapis, demografi usia, tabel klien kredit rendah, dan matriks kehadiran/kuota.", "Should"),
    ("Detail klien", "Demografi, paket kredit (renewal hanya melalui Finance), rutinitas mingguan, dan riwayat sesi.", "Must"),
    ("Discharge", "Alasan: Pindah domisili, Finansial, Bentrok jadwal, Ekspektasi tidak terpenuhi, Lulus/Graduate, Lainnya + catatan klinis. Status klien menjadi Discharged.", "Must"),
])

h("4.7 Keuangan, Paket & Kredit (FIN)", 2)
fr_table("FR-FIN", [
    ("Master paket", "Kelola paket: Nama, Jumlah sesi (kredit), Harga, Deskripsi. Contoh awal: Regular Therapist (10 sesi), Senior Therapist (10 sesi), Konsultasi (1 sesi), Sensori Intensif (15 sesi). Tarif final disediakan Therapedia.", "Must"),
    ("Terbitkan invoice", "Field: Klien, Paket, Nominal. Invoice tampil di Portal Orang Tua berstatus Belum Dibayar.", "Must"),
    ("Verifikasi transfer", "Finance melihat bukti transfer (zoom, rotasi, unduh) lalu Approve (invoice Lunas & kredit paket otomatis ditambahkan) atau Reject (kembali Belum Dibayar, disertai alasan).", "Must"),
    ("Renewal kredit", "Finance memperpanjang paket klien (Klien, Paket, Jumlah kredit, Nominal) dan sistem membuat invoice terkait.", "Must"),
    ("Arsip invoice", "Daftar seluruh invoice dengan status dan filter.", "Must"),
    ("Buku besar kredit", "Audit trail kredit per klien dengan jenis transaksi: Terpakai, Batal (dimaafkan), Batal (penalti), Renewal; berhalaman.", "Must"),
    ("Metode pembayaran", "Pembayaran melalui transfer bank manual (tanpa payment gateway) dengan unggah bukti.", "Must"),
])

h("4.8 Portal Terapis (THR)", 2)
fr_table("FR-THR", [
    ("Jadwal saya", "Tampilan minggu/hari khusus sesi milik terapis; filter klien dan status laporan (semua/belum/terisi); akses cepat ke profil, Drive, dan kuesioner klien. Terapis tidak dapat membatalkan/reschedule.", "Must"),
    ("Laporan sesi", "Formulir laporan: Activity, Note, Homework dengan template isian cepat. Laporan dianggap lengkap bila ketiga bagian terisi.", "Must"),
    ("Ringkasan laporan", "KPI: sesi selesai, laporan tertunda, laporan lengkap (%), jumlah klien; tab sesi selesai dan klien; filter klien, status, periode, pencarian.", "Must"),
    ("Riwayat laporan klien", "Panel riwayat seluruh laporan sesi seorang klien.", "Must"),
    ("Berkas klien", "Profil, rutinitas mingguan, riwayat sesi, dan hasil kuesioner (read-only).", "Must"),
    ("Cetak laporan klien", "Laporan perkembangan klien dapat dicetak / disimpan sebagai PDF.", "Should"),
])

h("4.9 Portal Orang Tua (PRT)", 2)
fr_table("FR-PRT", [
    ("Status tagihan", "Kartu invoice dengan 3 status: Menunggu Pembayaran (merah), Menunggu Verifikasi (kuning), Lunas Terverifikasi (hijau).", "Must"),
    ("Unggah bukti bayar", "Drag-and-drop / pilih berkas JPG, PNG, WEBP (maks. 25 MB, dikompresi otomatis) atau PDF (maks. 8 MB); dapat diunggah ulang dan dilihat kembali.", "Must"),
    ("Saldo kredit", "Menampilkan sisa kredit per paket.", "Must"),
    ("Link Google Drive", "Akses tautan folder dokumen anak.", "Must"),
    ("Riwayat sesi", "Daftar sesi yang telah selesai beserta laporan sesi dari terapis.", "Must"),
    ("Tampilan mobile", "Portal dioptimalkan untuk penggunaan di ponsel.", "Must"),
])

h("4.10 Dashboard Eksekutif (EXE)", 2)
fr_table("FR-EXE", [
    ("Dashboard revenue", "Periode minggu/bulan/kuartal/custom; filter cabang (Master). KPI klien, terapis, invoice, sesi, dan pendapatan terverifikasi. Grafik pendapatan & inquiry per cabang, pendapatan per paket; tabel beban terapis dan transaksi.", "Must"),
    ("Revenue cabang", "Branch Manager melihat dashboard revenue terbatas pada cabangnya.", "Must"),
    ("Performa cabang", "KPI total intake, admitted + % konversi, in-progress, discontinued, cabang terbaik; grafik perbandingan cabang dan tren 6 bulan; matriks perbandingan.", "Must"),
])

h("4.11 Administrasi Pengguna & RBAC (ADM)", 2)
fr_table("FR-ADM", [
    ("Daftar pengguna", "Tabel berhalaman dengan pencarian dan filter cabang.", "Must"),
    ("Tambah pengguna", "Field: Nama*, Email*, Peran, Cabang, Kata sandi awal; bila Terapis: Gelar, Spesialisasi, Bio. Terapis baru otomatis tersedia pada pilihan terapis di penjadwalan.", "Must"),
    ("Ubah & nonaktifkan", "Master dapat mengubah data pengguna, mereset kata sandi, dan menonaktifkan/menghapus akun dengan konfirmasi.", "Must"),
    ("Jam kerja terapis", "Pengaturan hari & jam kerja terapis sebagai dasar deteksi konflik jadwal.", "Must"),
    ("Matriks RBAC", "Toggle akses 10 modul × peran.", "Must"),
    ("Custom role", "Membuat, mengubah, dan menghapus peran baru (label, slug, badge, deskripsi, modul awal). Peran bawaan sistem tidak dapat dihapus. Pengguna dengan custom role dapat membuka seluruh modul yang diizinkan.", "Must"),
])

# ---------------------------------------------------------------------------
# 5. NON-FUNGSIONAL
# ---------------------------------------------------------------------------
h("5. Kebutuhan Non-Fungsional", 1)
table(["ID", "Kategori", "Kebutuhan"], [
    ("NFR-01", "Keamanan", "Kata sandi di-hash; seluruh komunikasi via HTTPS; hak akses diverifikasi di sisi server (bukan hanya tampilan)."),
    ("NFR-02", "Kerahasiaan data", "Data anak, hasil asesmen, dan laporan klinis hanya dapat diakses peran yang berwenang; tidak dibagikan ke pihak ketiga."),
    ("NFR-03", "Audit trail", "Aksi penting (perubahan status, verifikasi pembayaran, perubahan kredit, perubahan hak akses) tercatat: siapa, kapan, apa."),
    ("NFR-04", "Penyimpanan berkas", "Bukti transfer dan lampiran disimpan di server/penyimpanan berkas, bukan di browser."),
    ("NFR-05", "Performa", "Halaman utama dan dashboard termuat < 3 detik pada koneksi normal untuk volume data operasional klinik."),
    ("NFR-06", "Ketersediaan & backup", "Backup database otomatis harian; prosedur pemulihan terdokumentasi."),
    ("NFR-07", "Kompatibilitas", "Browser modern (Chrome, Edge, Safari, Firefox) versi terbaru; modul staf optimal di desktop/tablet, Portal Orang Tua & kuesioner optimal di ponsel."),
    ("NFR-08", "Bahasa & format", "Antarmuka berbahasa Indonesia (istilah teknis tertentu dapat berbahasa Inggris); mata uang Rupiah (IDR); format tanggal DD/MM/YYYY; zona waktu WIB."),
    ("NFR-09", "Usability", "Desain konsisten dengan identitas visual Therapedia; notifikasi aksi (toast) untuk setiap aksi berhasil/gagal."),
    ("NFR-10", "Infrastruktur", "Server VPS (min. 2 vCPU, 4 GB RAM) dan domain disediakan Therapedia; instalasi & konfigurasi oleh tim pengembang."),
], widths=[2, 3.4, 11.4])

# ---------------------------------------------------------------------------
# 6. ATURAN BISNIS
# ---------------------------------------------------------------------------
h("6. Aturan Bisnis", 1)
table(["ID", "Aturan"], [
    ("BR-01", "1 kredit = 1 sesi terapi. Kredit terpotong saat sesi terapi ditandai Selesai."),
    ("BR-02", "Sesi asesmen tidak memotong kredit paket terapi; penyelesaiannya mengubah status klien menjadi Assessment Done."),
    ("BR-03", "Pembatalan sesi: 3 pembatalan pertama per klien tidak memotong kredit (dimaafkan); pembatalan ke-4 dan seterusnya memotong 1 kredit (penalti)."),
    ("BR-04", "Klien dengan saldo kredit 0 berstatus Frozen; sesinya tetap terlihat dan ditandai pada kalender sampai dilakukan renewal."),
    ("BR-05", "Kredit paket bertambah otomatis hanya setelah Finance meng-Approve bukti transfer (atau melakukan renewal)."),
    ("BR-06", "Keputusan Admit menetapkan tanggal bergabung dan menjadikan calon klien sebagai klien aktif."),
    ("BR-07", "Keputusan Discontinue dan Discharge wajib disertai alasan."),
    ("BR-08", "Hanya Master dan Admin Schedule yang dapat menyelesaikan, membatalkan, atau me-reschedule sesi."),
    ("BR-09", "Terapis hanya melihat jadwal & klien yang ditanganinya, dan hanya dapat mengisi laporan sesi."),
    ("BR-10", "Kuesioner hanya dapat dikirim bila seluruh item telah dijawab; pengiriman otomatis mengubah status ke Assessment Done."),
    ("BR-11", "School Companion Profile hanya dapat dipilih bersama layanan BOT-A atau FOT-A."),
    ("BR-12", "Reschedule tidak dapat disimpan bila bentrok dengan jadwal terapis lain atau di luar jam kerja terapis."),
    ("BR-13", "Pengguna non-Master hanya dapat mengakses data cabangnya sendiri."),
], widths=[2, 14.8])

# ---------------------------------------------------------------------------
# 7. DI LUAR RUANG LINGKUP
# ---------------------------------------------------------------------------
h("7. Di Luar Ruang Lingkup", 1)
p("Hal-hal berikut tidak termasuk dalam rilis ini dan dapat diajukan sebagai Change Request / fase berikutnya:")
bullets([
    "Integrasi payment gateway / pembayaran online (pembayaran tetap melalui transfer manual).",
    "Fitur chat di dalam aplikasi; komunikasi menggunakan tautan WhatsApp (wa.me) tanpa WhatsApp Business API berbayar.",
    "Notifikasi otomatis via email, SMS, atau push notification.",
    "Aplikasi mobile native (Android/iOS); sistem berbasis web responsif.",
    "Dukungan multi-bahasa dan mode gelap (dark mode).",
    "Migrasi data historis dari sistem/pencatatan lama (kecuali disepakati terpisah).",
    "Integrasi dengan BPJS, asuransi, atau sistem akuntansi eksternal.",
])

# ---------------------------------------------------------------------------
# 8. ASUMSI & BATASAN
# ---------------------------------------------------------------------------
h("8. Asumsi & Batasan", 1)
bullets([
    "Prototipe antarmuka yang telah didemonstrasikan menjadi acuan tampilan & alur; data contoh pada prototipe tidak dimigrasikan ke sistem produksi.",
    "Butir instrumen asesmen, tarif paket terapi, data terapis, dan konten website disediakan oleh Therapedia.",
    "Server, domain, dan akun Google Drive klinik disediakan oleh Therapedia.",
    "Sistem dibangun sebagai aplikasi web dengan backend dan database terpusat; rancangan database detail disusun setelah dokumen ini disetujui.",
    "Perubahan kebutuhan setelah persetujuan dokumen ini diproses melalui mekanisme Change Request.",
])

# ---------------------------------------------------------------------------
# 9. POIN KONFIRMASI
# ---------------------------------------------------------------------------
h("9. Poin yang Memerlukan Konfirmasi Klien", 1)
p("Mohon Therapedia memberikan konfirmasi atas poin-poin berikut sebelum dokumen ditandatangani:")
table(["No", "Pertanyaan", "Jawaban Klien"], [
    ("1", "Format laporan sesi terapis: cukup Activity / Note / Homework, atau perlu format SOAP (Subjective, Objective, Assessment, Plan)?", ""),
    ("2", "Aturan pembatalan: apakah batas 3 kali pembatalan gratis berlaku per klien selamanya, per paket, atau per bulan?", ""),
    ("3", "Apakah sesi boleh tetap dijadwalkan untuk klien berstatus Frozen (kredit 0)?", ""),
    ("4", "Nomor WhatsApp resmi per cabang yang digunakan pada website, pengingat, dan ucapan ulang tahun.", ""),
    ("5", "Daftar final layanan, paket, dan tarif per cabang (apakah tarif berbeda antar cabang?).", ""),
    ("6", "Apakah Finance ditugaskan per cabang atau menangani seluruh cabang?", ""),
    ("7", "Apakah kode ASM perlu masa berlaku / hanya dapat digunakan sekali?", ""),
    ("8", "Apakah orang tua perlu dapat melihat hasil kuesioner asesmen di Portal Orang Tua?", ""),
    ("9", "Jam operasional & jam kerja terapis per cabang (untuk validasi jadwal).", ""),
], widths=[1, 9.5, 6.3])

# ---------------------------------------------------------------------------
# 10. PERSETUJUAN
# ---------------------------------------------------------------------------
h("10. Persetujuan", 1)
p("Dengan ditandatanganinya dokumen ini, kedua belah pihak menyetujui bahwa kebutuhan yang tercantum "
  "merupakan ruang lingkup final pengembangan sistem.")
sign = doc.add_table(rows=5, cols=2)
sign.style = "Table Grid"
sign.alignment = WD_TABLE_ALIGNMENT.CENTER
labels = [("Pihak Klien", "Pihak Pengembang"),
          ("Therapedia Center", VENDOR),
          ("\n\n\n\n", "\n\n\n\n"),
          ("Nama  : ______________________", "Nama  : ______________________"),
          ("Jabatan : ____________________\nTanggal : ____________________",
           "Jabatan : ____________________\nTanggal : ____________________")]
for ri, (a, b) in enumerate(labels):
    for ci, text in enumerate((a, b)):
        cell = sign.rows[ri].cells[ci]
        cell.text = text
        cell.width = Cm(8.4)
        if ri == 0:
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            set_cell_bg(cell, BRAND)
        if ri <= 1:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Minta Word memperbarui field (daftar isi) saat dokumen dibuka
settings = doc.settings.element
upd = OxmlElement("w:updateFields")
upd.set(qn("w:val"), "true")
settings.append(upd)

doc.save(OUT_FILE)
print(f"Dokumen tersimpan: {OUT_FILE}")
